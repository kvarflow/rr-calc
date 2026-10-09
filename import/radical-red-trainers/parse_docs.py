"""Reads a Radical Red 4.1 boss document (.xlsx) into JSON for build_trainers.js.

Usage: python3 parse_docs.py <boss document .xlsx> <output .json>

The documents' "Trainer Order" sheet lists every boss fight in order, each name linking to
its team block on another sheet. A block has the trainer in column C and up to six Pokémon
in columns E, J, O, T, Y and AD: species, level, (2 blank rows), nature, ability, item, four
moves, then a BASE STATS table whose EVs sit three columns to the right, and the Pokémon's
in-game Speed stat. Fights with several possible teams repeat the block, each under a note such as
"(!) IF RIVAL HAS SQUIRTLE" or "(!) RAIN TEAM"; a "BATTLE EFFECT: DOUBLES" line may sit
above a block too.
"""
import json
import re
import sys

import openpyxl

STATS = ['hp', 'at', 'df', 'sa', 'sd', 'sp']  # document order: HP, ATK, DEF, SPA, SPD, SPE
FIRST_MON_COLUMN, MON_COLUMN_STEP, HEADER_COLUMN = 5, 5, 3


def value(ws, row, col):
    return ws.cell(row=row, column=col).value


def is_text(v):
    return isinstance(v, str) and v.strip() != ''


def parse_team(ws, header_row):
    team = []
    for col in range(FIRST_MON_COLUMN, ws.max_column + 1, MON_COLUMN_STEP):
        species = value(ws, header_row, col)
        if not is_text(species):
            continue
        stats_row = next((r for r in range(header_row + 8, header_row + 16) if value(ws, r, col) == 'BASE STATS'), None)
        evs = {}
        if stats_row:
            for i, stat in enumerate(STATS):
                ev = value(ws, stats_row + 1 + i, col + 3)
                if isinstance(ev, (int, float)) and ev:
                    evs[stat] = int(ev)
        moves = [value(ws, header_row + k, col) for k in range(7, 11)]
        team.append({
            'species': species.strip(),
            'level': value(ws, header_row + 1, col),  # a number, or relative text like "Highest Lv -3"
            'nature': value(ws, header_row + 4, col),
            'ability': value(ws, header_row + 5, col),
            'item': value(ws, header_row + 6, col),
            'moves': [m.strip() for m in moves if is_text(m) and m.strip() != '-'],
            'evs': evs,
            'speed': value(ws, stats_row + 7, col + 3) if stats_row else None,
        })
    return team


def text_above(ws, header_row, prefixes):
    """A note in the Pokémon columns' first column a few rows above a block, e.g. "(!) RAIN TEAM"."""
    for row in range(header_row - 1, max(header_row - 7, 0), -1):
        text = value(ws, row, FIRST_MON_COLUMN)
        if is_text(text) and text.strip().upper().startswith(prefixes):
            return text.strip()
    return None


def parse_variants(ws, first_row):
    """Every block under one link: the first, plus later ones repeating its header."""
    header = value(ws, first_row, HEADER_COLUMN)
    variants = []
    for row in range(first_row, ws.max_row + 1):
        text = value(ws, row, HEADER_COLUMN)
        if not is_text(text):
            continue
        if text != header:
            break
        effect = text_above(ws, row, ('BATTLE EFFECT:',))
        # Team notes are "(!) RAIN TEAM" or, on some sheets, a bare "IF RIVAL HAS SQUIRTLE".
        variants.append({'note': text_above(ws, row, ('(!)', 'IF ')),
                         'battleEffect': effect.split(':', 1)[1].strip() if effect else None,
                         'team': parse_team(ws, row)})
    return header, variants


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    order_ws = wb['Trainer Order']
    order = []
    optional = False
    for row in order_ws.iter_rows():
        for c in row:
            if c.value == '(OPTIONAL)':
                optional = True  # applies to the next trainer listed
            if not (c.hyperlink and c.hyperlink.location):
                continue
            sheet, link_row = re.match(r"'?(.+?)'?![A-Z]+(\d+)", c.hyperlink.location).groups()
            ws = wb[sheet]
            # The link points at or just above the block's header cell.
            header_row = next(r for r in range(int(link_row), int(link_row) + 6) if is_text(value(ws, r, HEADER_COLUMN)))
            header, variants = parse_variants(ws, header_row)
            order.append({
                'name': c.value,
                'location': order_ws.cell(row=c.row + 1, column=c.column).value,
                'levelCap': order_ws.cell(row=c.row, column=c.column + 2).value,
                'optional': optional,
                'header': header,
                'variants': variants,
            })
            optional = False
    return order


if __name__ == '__main__':
    trainers = parse(sys.argv[1])
    with open(sys.argv[2], 'w') as f:
        json.dump(trainers, f, indent=1, default=str)
    print(f'{sys.argv[1]}: {len(trainers)} trainers, {sum(len(t["variants"]) for t in trainers)} teams')
